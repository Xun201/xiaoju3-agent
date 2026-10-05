# -*- coding: utf-8 -*-
"""AIC 技术方案 MD → docx(标题/段落/表格/图片逐元素,中文字体样式),
之后 soffice docx→PDF(Writer 滤镜,质量远好于 HTML 管线)。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

MD = r"F:\Orangepi_number3\docs\AIC2026_TECH_REPORT.md"
OUT = r"F:\Orangepi_number3\docs\AIC2026_TECH_REPORT.docx"
IMG_DIR = r"F:\Orangepi_number3\docs"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(11)
for s in ("Heading 1", "Heading 2", "Heading 3"):
    doc.styles[s].font.name = "Microsoft YaHei"
    doc.styles[s].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

lines = open(MD, encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip() or line.startswith("---"):
        i += 1
        continue
    # 图片(alt 可含方括号,如"[思考]/[计划]"——用非贪婪+右锚定)
    m = re.match(r"^!\[(.+?)\]\((img/[^)]+)\)\s*$", line.strip())
    if m:
        alt, rel = m.group(1), m.group(2)
        try:
            from PIL import Image as PILImage
            w, h = PILImage.open(IMG_DIR + "/" + rel).size
            doc.add_picture(IMG_DIR + "/" + rel, width=Cm(15.5))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(alt)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        except Exception as e:
            doc.add_paragraph(f"[图片缺失:{alt}({e})]")
        i += 1
        continue
    # 表格
    if line.strip().startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").replace(" ", "")) <= set("-:"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            rows.append(cells)
            i += 1
        rows.pop(1)  # 分隔行
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci in range(cols):
                cell = table.cell(ri, ci)
                txt = row[ci] if ci < len(row) else ""
                cell.text = re.sub(r"\*\*([^*]+)\*\*", r"\1", txt)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if ri == 0:
                            r.font.bold = True
        doc.add_paragraph()
        continue
    # 标题
    hm = re.match(r"^(#{1,3})\s+(.*)", line)
    if hm:
        level = len(hm.group(1))
        text = re.sub(r"\*\*([^*]+)\*\*", r"\1", hm.group(2))
        if level == 1:
            p = doc.add_heading(text, level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_heading(text, level=level)
        i += 1
        continue
    # 引用
    if line.startswith("> "):
        p = doc.add_paragraph(re.sub(r"\*\*([^*]+)\*\*", r"\1", line[2:]))
        p.paragraph_format.left_indent = Cm(0.6)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        i += 1
        continue
    # 列表（数字列表所见即所得：字面数字写进文本，不用 Word List Number
    # 自动编号——该样式跨列表续排，§1.2 的 1-4 会让 §3.4 显示成 5-8）
    lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
    if lm:
        prefix = lm.group(2) + " " if lm.group(2)[0].isdigit() else "• "
        text = re.sub(r"\*\*([^*]+)\*\*", r"\1", lm.group(3))
        style_name = "List Bullet" if lm.group(2)[0] in "-*" else None
        if style_name:
            try:
                doc.add_paragraph(text, style=style_name)
            except KeyError:
                doc.add_paragraph("• " + text)
        else:
            p = doc.add_paragraph(prefix + text)
            p.paragraph_format.left_indent = Cm(0.6)
        i += 1
        continue
    # 普通段落(行内 code/加粗处理:剥标记)
    text = re.sub(r"`([^`]+)`", r"\1", line)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    doc.add_paragraph(text)
    i += 1

doc.save(OUT)
print("docx 已生成:", OUT)
